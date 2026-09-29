"""Extract plain text from uploaded documents (.txt, .md, .docx, .pdf)."""
from __future__ import annotations

import io
from pathlib import Path

ALLOWED_EXT = {".txt", ".md", ".docx", ".pdf"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


class ExtractError(Exception):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name or ""
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ExtractError("unsupported")
    if uploaded_file.size and uploaded_file.size > MAX_UPLOAD_BYTES:
        raise ExtractError("too_large")

    data = uploaded_file.read()

    if ext in (".txt", ".md"):
        text = _decode(data)
    elif ext == ".docx":
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise ExtractError("missing_lib") from exc
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:  # .pdf
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise ExtractError("missing_lib") from exc
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)

    # normalise whitespace but keep paragraph breaks
    lines = [" ".join(line.split()) for line in text.splitlines()]
    cleaned = "\n".join(lines)
    while "\n\n\n" in cleaned:
        cleaned = cleaned.replace("\n\n\n", "\n\n")
    cleaned = cleaned.strip()
    if not cleaned:
        raise ExtractError("empty")
    return cleaned
